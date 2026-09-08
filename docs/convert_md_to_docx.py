"""Markdown -> docx 轻量转换 (支持标题/表格/列表/代码块/加粗/行内代码子集).

用法:
    python convert_md_to_docx.py <input.md> <output.docx>
"""
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor

SRC, DST = sys.argv[1], sys.argv[2]

doc = Document()
styles = doc.styles
for i in range(1, 4):
    styles[f"Heading {i}"].font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)
styles["Normal"].font.size = Pt(10.5)


def add_code(lines):
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def add_table(rows):
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.text = row[ci] if ci < len(row) else ""
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)


def rich(p, text):
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        else:
            p.add_run(tok)


with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

in_code = False
code_buf = []
i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith("```"):
        if in_code:
            add_code(code_buf)
            code_buf = []
            in_code = False
        else:
            in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(line)
        i += 1
        continue
    if not line.strip():
        i += 1
        continue
    if line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            parts = [c.strip() for c in lines[i].strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", p) for p in parts):
                rows.append(parts)
            i += 1
        add_table(rows)
        doc.add_paragraph()
        continue
    m = re.match(r"^(#{1,4})\s+(.*)", line)
    if m:
        p = doc.add_heading(level=min(len(m.group(1)), 4))
        rich(p, m.group(2))
        i += 1
        continue
    if re.fullmatch(r"-{3,}", line.strip()):
        i += 1
        continue
    m = re.match(r"^\s*[-*]\s+(.*)", line)
    if m:
        p = doc.add_paragraph(style="List Bullet")
        rich(p, m.group(1))
        i += 1
        continue
    m = re.match(r"^\s*\d+\.\s+(.*)", line)
    if m:
        p = doc.add_paragraph(style="List Number")
        rich(p, m.group(1))
        i += 1
        continue
    p = doc.add_paragraph()
    rich(p, line)
    i += 1

doc.save(DST)
print("saved:", DST)
